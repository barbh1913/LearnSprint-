"""The content-intelligence pipeline (FR2.8, FR2.10, ADR 0013).

The matcher is tested as a pure function; the flows through the real wiring
with the AI faked at the ai_extractor boundary. Uploads are real .pptx decks,
so the text extraction runs for real too.
"""

from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from features.content_topics.domain.material_matching import (
    ATTACH_EXISTING,
    CREATE_NEW,
    MaterialContent,
    TopicCandidate,
    normalize,
    rank_against_topics,
    terms,
)
from features.content_topics.infrastructure import ai_extractor
from main import app
from shared.config import settings

from .conftest import FakeBucket

client = TestClient(app)


def auth_headers(email: str = "student@example.com") -> dict[str, str]:
    response = client.post("/auth/register", json={"email": email, "password": "s3cret123"})
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def create_course(headers: dict[str, str]) -> dict[str, Any]:
    return client.post(
        "/courses",
        json={
            "name": "Algorithms",
            "year": 2,
            "semester": "A",
            "credits": 5,
            "examDate": (datetime.now() + timedelta(days=30)).isoformat(),
            "examType": "closed",
        },
        headers=headers,
    ).json()


def create_topic(headers: dict[str, str], course_id: str, name: str, description: str | None = None) -> dict[str, Any]:
    payload: dict[str, Any] = {"name": name}
    if description:
        payload["description"] = description
    return client.post(f"/courses/{course_id}/topics", json=payload, headers=headers).json()


def topics_of(headers: dict[str, str], course_id: str) -> dict[str, dict[str, Any]]:
    return {t["name"]: t for t in client.get(f"/courses/{course_id}/topics", headers=headers).json()}


def card_for(headers: dict[str, str], course_id: str, topic_id: str) -> dict[str, Any]:
    board = client.get(f"/board?courseId={course_id}", headers=headers).json()
    return next(card for card in board["cards"] if card["topicId"] == topic_id)


def make_pptx(*slide_titles: str) -> bytes:
    presentation = Presentation()
    layout = presentation.slide_layouts[0]
    for title in slide_titles:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def analyze(headers: dict[str, str], course_id: str, name: str, *slides: str, kind: str = "materials"):
    return client.post(
        f"/courses/{course_id}/{kind}/analyze",
        files={"file": (name, make_pptx(*slides), "application/vnd.openxmlformats")},
        headers=headers,
    )


def content(title: str, *key_points: str, topics: tuple[str, ...] = ()) -> MaterialContent:
    return MaterialContent(
        title=title,
        summary=None,
        key_points=tuple(key_points),
        topics=topics,
        estimated_minutes=None,
        language="en",
    )


# --- The matcher ------------------------------------------------------------------


class TestMatcher:
    graphs = TopicCandidate("t-graphs", "Graph Algorithms", "BFS, DFS and shortest paths")
    trees = TopicCandidate("t-trees", "Binary Trees", "Traversal, balancing, heaps")
    hashing = TopicCandidate("t-hash", "Hashing", None)

    def test_strong_match_when_the_material_names_the_topic(self) -> None:
        match = rank_against_topics(content("Lecture 5", "BFS", "DFS", topics=("Graph Algorithms",)), [self.trees, self.graphs])

        assert match.decision == ATTACH_EXISTING
        assert match.best is not None and match.best.topic_id == "t-graphs"
        assert match.best.confidence >= 0.9
        assert "names 'Graph Algorithms'" in match.reason

    def test_key_terms_carry_a_match_when_titles_differ(self) -> None:
        match = rank_against_topics(content("Exercise sheet 3", "BFS", "DFS", "shortest paths"), [self.trees, self.graphs, self.hashing])

        assert match.decision == ATTACH_EXISTING
        assert match.best is not None and match.best.topic_id == "t-graphs"
        assert "shares" in match.reason and "bfs" in match.reason

    def test_no_meaningful_match_recommends_a_new_topic_and_says_why(self) -> None:
        match = rank_against_topics(content("Dynamic Programming", "memoization", "tabulation"), [self.trees, self.graphs])

        assert match.decision == CREATE_NEW
        assert match.best is None
        assert match.suggested_title == "Dynamic Programming"
        assert "No existing topic covers this material" in match.reason

    def test_alternatives_are_ranked_and_capped(self) -> None:
        heaps = TopicCandidate("t-heaps", "Heaps", "priority queues and binary trees")
        match = rank_against_topics(content("Binary Trees", "traversal", "heaps"), [self.graphs, heaps, self.trees, self.hashing])

        assert match.best is not None and match.best.topic_id == "t-trees"
        assert [alt.topic_id for alt in match.alternatives] == ["t-heaps"]
        assert all(alt.confidence <= match.best.confidence for alt in match.alternatives)

    def test_a_course_without_topics_creates_the_first_one(self) -> None:
        match = rank_against_topics(content("Sorting"), [])

        assert (match.decision, match.suggested_title) == (CREATE_NEW, "Sorting")
        assert "no topics yet" in match.reason

    def test_hebrew_and_punctuation_normalise_the_same_way(self) -> None:
        hebrew = TopicCandidate("t-he", "עצים בינאריים", "מעבר על עץ")
        match = rank_against_topics(content("שיעור 4: עצים בינאריים!", "מעבר על עץ"), [self.graphs, hebrew])

        assert match.best is not None and match.best.topic_id == "t-he"
        assert normalize("Graph-Algorithms, BFS!") == "graph algorithms bfs"
        assert terms("The BFS and the DFS", "of graphs") == {"bfs", "dfs", "graphs"}

    def test_ties_break_deterministically(self) -> None:
        # Two topics that score identically: name first, then id - never input order.
        later = TopicCandidate("t-b", "Sorting", None)
        earlier = TopicCandidate("t-a", "Sorting", None)
        match = rank_against_topics(content("Sorting", "quicksort"), [later, earlier])

        assert match.best is not None and match.best.topic_id == "t-a"
        assert rank_against_topics(content("Sorting", "quicksort"), [earlier, later]).best.topic_id == "t-a"


# --- Single-material flow ----------------------------------------------------------


class TestMaterialAnalysis:
    def test_analysis_stores_the_file_and_recommends_without_changing_the_course(
        self, fake_storage: FakeBucket
    ) -> None:
        headers = auth_headers()
        course = create_course(headers)
        create_topic(headers, course["id"], "Graph Algorithms", "BFS, DFS and shortest paths")
        create_topic(headers, course["id"], "Binary Trees")

        response = analyze(headers, course["id"], "lecture5.pptx", "Graph Algorithms", "BFS", "DFS", "Dijkstra")

        assert response.status_code == 201, response.text
        body = response.json()
        assert body["analysedBy"] == "heuristic"
        assert body["content"]["title"] == "Graph Algorithms"
        assert "BFS" in body["content"]["keyPoints"]
        assert body["recommendation"]["decision"] == "attach_existing"
        assert body["recommendation"]["topicName"] == "Graph Algorithms"
        assert body["recommendation"]["confidence"] >= 0.9
        assert body["recommendation"]["reason"]
        # Stored, but not filed anywhere yet - and no topic was created.
        assert len(fake_storage.objects) == 1
        assert len(topics_of(headers, course["id"])) == 2
        graphs_id = topics_of(headers, course["id"])["Graph Algorithms"]["id"]
        assert client.get(f"/courses/{course['id']}/topics/{graphs_id}/materials", headers=headers).json() == []

    def test_one_file_is_one_topic_with_at_most_five_key_points(self) -> None:
        # A deck with many headings must not become many topics: the first
        # heading names the unit, the rest are its key points, capped at five.
        headers = auth_headers()
        course = create_course(headers)
        headings = ["Graph Algorithms"] + [f"{i}. Section {i}" for i in range(1, 10)]

        body = analyze(headers, course["id"], "lecture5.pptx", *headings).json()

        assert body["content"]["title"] == "Graph Algorithms"
        assert body["content"]["keyPoints"] == [f"Section {i}" for i in range(1, 6)]
        assert body["recommendation"]["decision"] == "create_new"
        assert body["recommendation"]["suggestedTitle"] == "Graph Algorithms"
        assert len(topics_of(headers, course["id"])) == 0

    def test_attaching_to_the_recommended_topic_files_the_material_and_refreshes_the_description(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        graphs = create_topic(headers, course["id"], "Graph Algorithms", "old description")
        analysis = analyze(headers, course["id"], "lecture5.pptx", "Graph Algorithms", "BFS", "DFS").json()

        response = client.post(
            f"/courses/{course['id']}/materials/{analysis['materialId']}/confirm",
            json={"decision": "attach", "topicId": analysis["recommendation"]["topicId"]},
            headers=headers,
        )

        assert response.status_code == 200, response.text
        assert response.json() == {
            "materialId": analysis["materialId"],
            "topicId": graphs["id"],
            "topicName": "Graph Algorithms",
            "created": False,
            "alreadyConfirmed": False,
        }
        listed = client.get(f"/courses/{course['id']}/topics/{graphs['id']}/materials", headers=headers).json()
        assert [m["fileName"] for m in listed] == ["lecture5.pptx"]
        description = topics_of(headers, course["id"])["Graph Algorithms"]["description"]
        assert "Key points:" in description and "- BFS" in description

    def test_attaching_never_touches_existing_progress(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        graphs = create_topic(headers, course["id"], "Graph Algorithms")
        read = card_for(headers, course["id"], graphs["id"])["actions"][0]
        client.patch(f"/actions/{read['id']}/progress?courseId={course['id']}", json={"isDone": True}, headers=headers)
        client.patch(f"/topics/{graphs['id']}/progress?courseId={course['id']}", json={"status": "in_progress"}, headers=headers)
        analysis = analyze(headers, course["id"], "more.pptx", "Graph Algorithms", "Dijkstra").json()

        client.post(
            f"/courses/{course['id']}/materials/{analysis['materialId']}/confirm",
            json={"decision": "attach", "topicId": graphs["id"]},
            headers=headers,
        )

        card = card_for(headers, course["id"], graphs["id"])
        assert card["actionsDone"] == 1
        assert card["status"] == "in_progress"
        assert [a["title"] for a in card["actions"]] == ["Read", "Summarize", "Quiz"]

    def test_choosing_another_topic_is_respected(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        create_topic(headers, course["id"], "Graph Algorithms", "BFS DFS")
        trees = create_topic(headers, course["id"], "Binary Trees")
        analysis = analyze(headers, course["id"], "lecture5.pptx", "Graph Algorithms", "BFS").json()
        assert analysis["recommendation"]["topicName"] == "Graph Algorithms"

        response = client.post(
            f"/courses/{course['id']}/materials/{analysis['materialId']}/confirm",
            json={"decision": "attach", "topicId": trees["id"]},
            headers=headers,
        )

        assert response.json()["topicName"] == "Binary Trees"

    def test_creating_a_new_topic_uses_the_estimate_and_the_editable_title(self, monkeypatch: pytest.MonkeyPatch) -> None:
        headers = auth_headers()
        course = create_course(headers)
        create_topic(headers, course["id"], "Binary Trees")
        monkeypatch.setattr(settings, "system_anthropic_api_key", "sk-test")
        monkeypatch.setattr(
            ai_extractor,
            "understand_material",
            lambda lines, api_key: ai_extractor.MaterialUnderstanding(
                title="Dynamic Programming",
                summary="How overlapping subproblems are solved once and reused.",
                key_points=["Memoization", "Tabulation", "Optimal substructure"],
                topics=["Dynamic Programming"],
                estimated_minutes=180,
                language="en",
            ),
        )
        analysis = analyze(headers, course["id"], "lecture9.pptx", "whatever").json()
        assert analysis["analysedBy"] == "ai"
        assert analysis["recommendation"]["decision"] == "create_new"
        assert analysis["recommendation"]["suggestedTitle"] == "Dynamic Programming"

        response = client.post(
            f"/courses/{course['id']}/materials/{analysis['materialId']}/confirm",
            json={"decision": "create", "title": "DP (Dynamic Programming)"},
            headers=headers,
        )

        assert response.status_code == 200, response.text
        body = response.json()
        assert body["created"] is True and body["topicName"] == "DP (Dynamic Programming)"
        card = card_for(headers, course["id"], body["topicId"])
        assert card["totalMinutes"] == 180
        assert [a["title"] for a in card["actions"]] == ["Read", "Summarize", "Quiz"]
        assert card["description"].startswith("How overlapping subproblems")
        assert "- Memoization" in card["description"]

    def test_confirming_twice_returns_the_first_result(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        analysis = analyze(headers, course["id"], "lecture1.pptx", "Sorting", "Quicksort").json()
        url = f"/courses/{course['id']}/materials/{analysis['materialId']}/confirm"

        first = client.post(url, json={"decision": "create"}, headers=headers).json()
        second = client.post(url, json={"decision": "create"}, headers=headers).json()

        assert first["created"] is True and first["alreadyConfirmed"] is False
        assert second == {**first, "created": False, "alreadyConfirmed": True}
        assert len(topics_of(headers, course["id"])) == 1

    def test_creating_a_topic_that_already_exists_reuses_it(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        existing = create_topic(headers, course["id"], "Sorting")
        analysis = analyze(headers, course["id"], "lecture1.pptx", "Quicksort details").json()

        response = client.post(
            f"/courses/{course['id']}/materials/{analysis['materialId']}/confirm",
            json={"decision": "create", "title": "  sorting "},
            headers=headers,
        )

        assert response.json()["topicId"] == existing["id"]
        assert response.json()["created"] is False
        assert len(topics_of(headers, course["id"])) == 1

    def test_a_topic_deleted_after_analysis_is_reported_not_guessed(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        graphs = create_topic(headers, course["id"], "Graph Algorithms")
        analysis = analyze(headers, course["id"], "lecture5.pptx", "Graph Algorithms").json()
        client.delete(f"/courses/{course['id']}/topics/{graphs['id']}", headers=headers)

        response = client.post(
            f"/courses/{course['id']}/materials/{analysis['materialId']}/confirm",
            json={"decision": "attach", "topicId": graphs["id"]},
            headers=headers,
        )

        assert response.status_code == 404
        assert client.get(f"/courses/{course['id']}/topics", headers=headers).json() == []

    def test_ai_failure_falls_back_to_the_heuristic_and_keeps_the_file(
        self, monkeypatch: pytest.MonkeyPatch, fake_storage: FakeBucket
    ) -> None:
        headers = auth_headers()
        course = create_course(headers)
        monkeypatch.setattr(settings, "system_anthropic_api_key", "sk-test")

        def boom(lines, api_key):
            raise ai_extractor.AiAnalysisUnavailable("rate limited")

        monkeypatch.setattr(ai_extractor, "understand_material", boom)

        response = analyze(headers, course["id"], "lecture5.pptx", "Graph Algorithms", "BFS")

        assert response.status_code == 201
        body = response.json()
        assert body["analysedBy"] == "heuristic"
        assert "rate limited" in body["note"]
        assert body["content"]["title"] == "Graph Algorithms"
        assert len(fake_storage.objects) == 1

    def test_hebrew_material_is_understood_in_hebrew(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        create_topic(headers, course["id"], "עצים בינאריים")

        body = analyze(headers, course["id"], "shiur4.pptx", "עצים בינאריים", "מעבר על עץ").json()

        assert body["content"]["language"] == "he"
        assert body["recommendation"]["topicName"] == "עצים בינאריים"

    def test_unreadable_and_unsupported_files_are_refused_and_not_kept(self, fake_storage: FakeBucket) -> None:
        headers = auth_headers()
        course = create_course(headers)

        unsupported = client.post(
            f"/courses/{course['id']}/materials/analyze",
            files={"file": ("notes.txt", b"plain text", "text/plain")},
            headers=headers,
        )
        corrupt = client.post(
            f"/courses/{course['id']}/materials/analyze",
            files={"file": ("broken.pdf", b"%PDF-1.4 not really", "application/pdf")},
            headers=headers,
        )

        assert unsupported.status_code == 400
        assert corrupt.status_code == 400
        assert fake_storage.objects == {}

    def test_only_members_can_analyse_and_only_the_uploader_can_confirm(self) -> None:
        owner = auth_headers("owner@example.com")
        course = create_course(owner)
        stranger = auth_headers("stranger@example.com")
        member = auth_headers("peer@example.com")
        client.post(f"/courses/{course['id']}/members", json={"email": "peer@example.com"}, headers=owner)
        analysis = analyze(owner, course["id"], "lecture1.pptx", "Sorting").json()

        assert analyze(stranger, course["id"], "x.pptx", "Sorting").status_code == 403
        response = client.post(
            f"/courses/{course['id']}/materials/{analysis['materialId']}/confirm",
            json={"decision": "create"},
            headers=member,
        )
        assert response.status_code == 403


# --- Syllabus flow -------------------------------------------------------------------


def fake_lectures(*titles: str):
    return lambda lines, api_key: [
        ai_extractor.LectureProposal(
            title=title,
            summary=f"What {title} covers.",
            key_points=[f"{title} point 1", f"{title} point 2"],
            estimated_minutes=120,
            language="en",
        )
        for title in titles
    ]


class TestSyllabusAnalysis:
    def test_proposes_lecture_level_topics_with_matches_and_creates_nothing_yet(self, monkeypatch: pytest.MonkeyPatch) -> None:
        headers = auth_headers()
        course = create_course(headers)
        graphs = create_topic(headers, course["id"], "Graph Algorithms")
        monkeypatch.setattr(settings, "system_anthropic_api_key", "sk-test")
        monkeypatch.setattr(ai_extractor, "propose_lecture_structure", fake_lectures("Sorting", "Graph Algorithms", "Hashing"))

        response = analyze(headers, course["id"], "syllabus.pptx", "Course outline", kind="syllabus")

        assert response.status_code == 201, response.text
        body = response.json()
        assert body["analysedBy"] == "ai"
        assert [p["title"] for p in body["proposals"]] == ["Sorting", "Graph Algorithms", "Hashing"]
        assert body["proposals"][0]["match"]["decision"] == "create_new"
        assert body["proposals"][1]["match"] == {
            **body["proposals"][1]["match"],
            "decision": "attach_existing",
            "topicId": graphs["id"],
        }
        assert body["proposals"][1]["estimatedMinutes"] == 120
        assert len(topics_of(headers, course["id"])) == 1

    def test_confirming_creates_attaches_and_skips_as_reviewed(self, monkeypatch: pytest.MonkeyPatch) -> None:
        headers = auth_headers()
        course = create_course(headers)
        graphs = create_topic(headers, course["id"], "Graph Algorithms")
        monkeypatch.setattr(settings, "system_anthropic_api_key", "sk-test")
        monkeypatch.setattr(ai_extractor, "propose_lecture_structure", fake_lectures("Sorting", "Graph Algorithms", "Hashing"))
        analysis = analyze(headers, course["id"], "syllabus.pptx", "outline", kind="syllabus").json()

        response = client.post(
            f"/courses/{course['id']}/syllabus/confirm",
            json={
                "materialId": analysis["materialId"],
                "items": [
                    {"index": 0, "decision": "create", "title": "Sorting algorithms"},
                    {"index": 1, "decision": "attach", "topicId": graphs["id"]},
                    {"index": 2, "decision": "skip"},
                ],
            },
            headers=headers,
        )

        assert response.status_code == 200, response.text
        body = response.json()
        assert [t["name"] for t in body["created"]] == ["Sorting algorithms"]
        assert [t["topicId"] for t in body["attached"]] == [graphs["id"]]
        assert body["skipped"] == 1
        assert body["alreadyConfirmed"] is False
        names = topics_of(headers, course["id"])
        assert set(names) == {"Graph Algorithms", "Sorting algorithms"}
        assert "Sorting point 1" in names["Sorting algorithms"]["description"]
        assert card_for(headers, course["id"], names["Sorting algorithms"]["id"])["totalMinutes"] == 120
        assert "Graph Algorithms point 1" in names["Graph Algorithms"]["description"]

    def test_confirming_twice_creates_nothing_more(self, monkeypatch: pytest.MonkeyPatch) -> None:
        headers = auth_headers()
        course = create_course(headers)
        monkeypatch.setattr(settings, "system_anthropic_api_key", "sk-test")
        monkeypatch.setattr(ai_extractor, "propose_lecture_structure", fake_lectures("Sorting", "Hashing"))
        analysis = analyze(headers, course["id"], "syllabus.pptx", "outline", kind="syllabus").json()
        payload = {
            "materialId": analysis["materialId"],
            "items": [{"index": 0, "decision": "create"}, {"index": 1, "decision": "create"}],
        }

        first = client.post(f"/courses/{course['id']}/syllabus/confirm", json=payload, headers=headers).json()
        second = client.post(f"/courses/{course['id']}/syllabus/confirm", json=payload, headers=headers).json()

        assert second == {**first, "alreadyConfirmed": True}
        assert len(topics_of(headers, course["id"])) == 2

    def test_without_ai_each_heading_is_a_proposal(self) -> None:
        headers = auth_headers()
        course = create_course(headers)

        body = analyze(headers, course["id"], "syllabus.pptx", "Week 1: Sorting", "Week 2: Hashing", kind="syllabus").json()

        assert body["analysedBy"] == "heuristic"
        # Headings are kept as written - the heuristic never invents structure.
        assert [p["title"] for p in body["proposals"]] == ["Week 1: Sorting", "Week 2: Hashing"]
        assert all(p["estimatedMinutes"] is None for p in body["proposals"])

    def test_a_proposal_outside_the_analysis_is_refused(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        analysis = analyze(headers, course["id"], "syllabus.pptx", "Week 1: Sorting", kind="syllabus").json()

        response = client.post(
            f"/courses/{course['id']}/syllabus/confirm",
            json={"materialId": analysis["materialId"], "items": [{"index": 7, "decision": "create"}]},
            headers=headers,
        )

        assert response.status_code == 422

    def test_a_syllabus_cannot_be_confirmed_as_a_single_material(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        analysis = analyze(headers, course["id"], "syllabus.pptx", "Week 1: Sorting", kind="syllabus").json()

        response = client.post(
            f"/courses/{course['id']}/materials/{analysis['materialId']}/confirm",
            json={"decision": "create"},
            headers=headers,
        )

        assert response.status_code == 400

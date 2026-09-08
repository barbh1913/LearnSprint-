"""End-to-end API tests covering the golden path.

Register -> create course -> add topics -> tick actions -> rate mastery ->
board updates -> schedule generates -> grades average. DynamoDB is faked in
conftest.py, so this runs offline.
"""

from __future__ import annotations

from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient

from main import app

client = TestClient(app)


def auth_headers(email: str = "student@example.com") -> dict[str, str]:
    response = client.post("/auth/register", json={"email": email, "password": "s3cret123"})
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def create_course(headers: dict[str, str], **overrides) -> dict:
    payload = {
        "name": "Data Structures",
        "year": 2,
        "semester": "A",
        "credits": 5,
        "examDate": (datetime.now() + timedelta(days=21)).isoformat(),
        "examType": "closed",
    }
    payload.update(overrides)
    return client.post("/courses", json=payload, headers=headers).json()


class TestCourses:
    def test_create_and_list_course(self) -> None:
        headers = auth_headers()
        created = create_course(headers)

        courses = client.get("/courses", headers=headers).json()

        assert created["name"] == "Data Structures"
        assert [course["id"] for course in courses] == [created["id"]]

    def test_courses_require_authentication(self) -> None:
        assert client.get("/courses").status_code == 401

    def test_another_user_cannot_see_my_course(self) -> None:
        owner = auth_headers("owner@example.com")
        course = create_course(owner)
        stranger = auth_headers("stranger@example.com")

        response = client.get(f"/courses/{course['id']}", headers=stranger)

        assert response.status_code == 403

    def test_delete_course_removes_it_from_the_list(self) -> None:
        headers = auth_headers()
        course = create_course(headers)

        assert client.delete(f"/courses/{course['id']}", headers=headers).status_code == 204
        assert client.get("/courses", headers=headers).json() == []

    def test_owner_can_update_the_course(self) -> None:
        headers = auth_headers()
        course = create_course(headers)

        response = client.patch(
            f"/courses/{course['id']}", json={"name": "Renamed"}, headers=headers
        )

        assert response.status_code == 200
        assert response.json()["name"] == "Renamed"

    def test_a_member_cannot_update_the_course(self) -> None:
        # Unlike topics, course fields like the exam date drive every member's
        # own schedule - only the owner changes them.
        owner = auth_headers("owner@example.com")
        peer = auth_headers("peer@example.com")
        course = create_course(owner)
        client.post(
            f"/courses/{course['id']}/members", json={"email": "peer@example.com"}, headers=owner
        )

        response = client.patch(
            f"/courses/{course['id']}", json={"name": "Hijacked"}, headers=peer
        )

        assert response.status_code == 403
        assert client.get(f"/courses/{course['id']}", headers=owner).json()["name"] != "Hijacked"

    def test_deleting_a_course_cleans_up_every_members_progress(self, fake_dynamo) -> None:
        # Progress rows are private, under each member's own USER# partition -
        # deleting the shared course must not leave them orphaned there forever.
        owner = auth_headers("owner@example.com")
        peer = auth_headers("peer@example.com")
        course = create_course(owner)
        client.post(
            f"/courses/{course['id']}/members", json={"email": "peer@example.com"}, headers=owner
        )
        topic = client.post(
            f"/courses/{course['id']}/topics", json={"name": "Shared topic"}, headers=owner
        ).json()
        peer_id = client.get("/auth/me", headers=peer).json()["id"]

        client.patch(
            f"/topics/{topic['id']}/progress?courseId={course['id']}",
            json={"masteryLevel": 3},
            headers=peer,
        )
        board = client.get(f"/board?courseId={course['id']}", headers=peer).json()
        action_id = board["cards"][0]["actions"][0]["id"]
        client.patch(
            f"/actions/{action_id}/progress?courseId={course['id']}",
            json={"isDone": True},
            headers=peer,
        )

        client.delete(f"/courses/{course['id']}", headers=owner)

        remaining = fake_dynamo.query(f"USER#{peer_id}")
        assert not any(topic["id"] in item["SK"] for item in remaining)
        assert not any(action_id in item["SK"] for item in remaining)


class TestGrades:
    def test_weighted_average_favours_heavier_courses(self) -> None:
        headers = auth_headers()
        big = create_course(headers, name="Big", credits=10)
        small = create_course(headers, name="Small", credits=2)

        client.put(f"/courses/{big['id']}/grade", json={"finalGrade": 90}, headers=headers)
        client.put(f"/courses/{small['id']}/grade", json={"finalGrade": 60}, headers=headers)

        grades = client.get("/grades", headers=headers).json()

        # (90*10 + 60*2) / 12 = 85, not the naive 75.
        assert grades["overall"]["average"] == 85.0

    def test_average_is_none_before_any_grade_is_entered(self) -> None:
        headers = auth_headers()
        create_course(headers)

        grades = client.get("/grades", headers=headers).json()

        assert grades["overall"]["average"] is None


class TestAiSettings:
    def test_defaults_before_anything_is_saved(self) -> None:
        settings = client.get("/ai-settings", headers=auth_headers()).json()

        assert settings == {"aiEnabled": False, "hasApiKey": False, "keyHint": None}

    def test_enabling_without_a_key_is_rejected(self) -> None:
        response = client.put(
            "/ai-settings", json={"aiEnabled": True}, headers=auth_headers()
        )

        assert response.status_code == 422

    def test_enabling_with_a_key_the_provider_rejects_fails(self, monkeypatch) -> None:
        monkeypatch.setattr(
            "features.academic_profile.presentation.router.verify_api_key", lambda key: False
        )

        response = client.put(
            "/ai-settings",
            json={"aiEnabled": True, "apiKey": "sk-ant-not-actually-valid"},
            headers=auth_headers(),
        )

        assert response.status_code == 422

    def test_enabling_with_a_valid_key_succeeds_and_never_echoes_it(self, monkeypatch) -> None:
        monkeypatch.setattr(
            "features.academic_profile.presentation.router.verify_api_key", lambda key: True
        )
        headers = auth_headers()

        response = client.put(
            "/ai-settings",
            json={"aiEnabled": True, "apiKey": "sk-ant-abcdef1234"},
            headers=headers,
        )

        body = response.json()
        assert response.status_code == 200
        assert body["aiEnabled"] is True
        assert body["hasApiKey"] is True
        assert body["keyHint"] == "...1234"
        assert "apiKey" not in body

    def test_toggling_off_keeps_the_saved_key(self, monkeypatch) -> None:
        monkeypatch.setattr(
            "features.academic_profile.presentation.router.verify_api_key", lambda key: True
        )
        headers = auth_headers()
        client.put(
            "/ai-settings",
            json={"aiEnabled": True, "apiKey": "sk-ant-abcdef1234"},
            headers=headers,
        )

        # No apiKey in this request - the frontend never holds the real key.
        response = client.put("/ai-settings", json={"aiEnabled": False}, headers=headers)

        body = response.json()
        assert body["aiEnabled"] is False
        assert body["hasApiKey"] is True
        assert body["keyHint"] == "...1234"


class TestTopicsAndBoard:
    def test_new_topic_gets_three_learning_actions(self) -> None:
        headers = auth_headers()
        course = create_course(headers)

        client.post(
            f"/courses/{course['id']}/topics", json={"name": "Binary Trees"}, headers=headers
        )
        board = client.get(f"/board?courseId={course['id']}", headers=headers).json()

        card = board["cards"][0]
        assert {action["type"] for action in card["actions"]} == {"read", "summarize", "quiz"}

    def test_topic_starts_in_backlog(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        client.post(f"/courses/{course['id']}/topics", json={"name": "Graphs"}, headers=headers)

        board = client.get(f"/board?courseId={course['id']}", headers=headers).json()

        assert board["cards"][0]["status"] == "backlog"

    def test_completing_one_action_moves_topic_to_in_progress(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        client.post(f"/courses/{course['id']}/topics", json={"name": "Graphs"}, headers=headers)
        board = client.get(f"/board?courseId={course['id']}", headers=headers).json()
        action_id = board["cards"][0]["actions"][0]["id"]

        client.patch(
            f"/actions/{action_id}/progress?courseId={course['id']}",
            json={"isDone": True},
            headers=headers,
        )
        board = client.get(f"/board?courseId={course['id']}", headers=headers).json()

        assert board["cards"][0]["status"] == "in_progress"

    def test_all_actions_done_prompts_for_mastery(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        card = _add_topic_and_finish_actions(headers, course["id"], "Sorting")

        assert card["needsMasteryRating"] is True
        assert card["status"] == "in_progress"

    def test_low_mastery_sends_topic_to_needs_review(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        card = _add_topic_and_finish_actions(headers, course["id"], "Sorting")

        client.patch(
            f"/topics/{card['topicId']}/progress?courseId={course['id']}",
            json={"masteryLevel": 2},
            headers=headers,
        )
        board = client.get(f"/board?courseId={course['id']}", headers=headers).json()

        assert board["cards"][0]["status"] == "needs_review"

    def test_high_mastery_completes_the_topic(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        card = _add_topic_and_finish_actions(headers, course["id"], "Sorting")

        client.patch(
            f"/topics/{card['topicId']}/progress?courseId={course['id']}",
            json={"masteryLevel": 5},
            headers=headers,
        )
        board = client.get(f"/board?courseId={course['id']}", headers=headers).json()

        assert board["cards"][0]["status"] == "done"

    def test_dragging_a_card_sets_its_status(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        topic = client.post(
            f"/courses/{course['id']}/topics", json={"name": "Heaps"}, headers=headers
        ).json()

        client.patch(
            f"/topics/{topic['id']}/progress?courseId={course['id']}",
            json={"status": "todo"},
            headers=headers,
        )
        board = client.get(f"/board?courseId={course['id']}", headers=headers).json()

        assert board["cards"][0]["status"] == "todo"

    def test_unknown_status_is_rejected(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        topic = client.post(
            f"/courses/{course['id']}/topics", json={"name": "Heaps"}, headers=headers
        ).json()

        response = client.patch(
            f"/topics/{topic['id']}/progress?courseId={course['id']}",
            json={"status": "not-a-real-status"},
            headers=headers,
        )

        assert response.status_code == 422

    def test_global_board_spans_every_course(self) -> None:
        headers = auth_headers()
        first = create_course(headers, name="Algorithms")
        second = create_course(headers, name="Databases")
        client.post(f"/courses/{first['id']}/topics", json={"name": "Sorting"}, headers=headers)
        client.post(f"/courses/{second['id']}/topics", json={"name": "Indexes"}, headers=headers)

        board = client.get("/board", headers=headers).json()

        assert board["totalTopics"] == 2
        assert {card["courseName"] for card in board["cards"]} == {"Algorithms", "Databases"}

    def test_renaming_and_prioritising_a_topic(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        topic = client.post(
            f"/courses/{course['id']}/topics", json={"name": "Old name"}, headers=headers
        ).json()

        updated = client.patch(
            f"/courses/{course['id']}/topics/{topic['id']}",
            json={"name": "New name", "isPriority": True},
            headers=headers,
        ).json()

        assert updated["name"] == "New name"
        assert updated["isPriority"] is True


class TestSchedule:
    def test_schedule_covers_the_topics(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        client.post(f"/courses/{course['id']}/topics", json={"name": "Recursion"}, headers=headers)

        schedule = client.get(f"/courses/{course['id']}/schedule", headers=headers).json()

        assert schedule["feasible"] is True
        assert len(schedule["blocks"]) > 0

    def test_schedule_needs_an_exam_date(self) -> None:
        headers = auth_headers()
        course = create_course(headers, examDate=None)

        response = client.get(f"/courses/{course['id']}/schedule", headers=headers)

        assert response.status_code == 400

    def test_blocked_hours_are_respected(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        client.post(f"/courses/{course['id']}/topics", json={"name": "Recursion"}, headers=headers)
        client.put(
            "/constraints",
            json={
                "blockedSlots": [{"day": 0, "startTime": "15:00", "endTime": "23:00"}],
                "timePreference": "evening",
            },
            headers=headers,
        )

        schedule = client.get(f"/courses/{course['id']}/schedule", headers=headers).json()

        mondays = [
            block
            for block in schedule["blocks"]
            if datetime.fromisoformat(block["start"]).weekday() == 0
        ]
        assert mondays == []

    def test_imminent_exam_switches_to_emergency_mode(self) -> None:
        headers = auth_headers()
        course = create_course(
            headers, examDate=(datetime.now() + timedelta(hours=6)).isoformat()
        )
        for name in ("A", "B", "C"):
            client.post(f"/courses/{course['id']}/topics", json={"name": name}, headers=headers)

        schedule = client.get(f"/courses/{course['id']}/schedule", headers=headers).json()

        assert schedule["feasible"] is False or schedule["isEmergencyMode"] is True


class TestVelocity:
    def test_velocity_counts_completed_actions(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        _add_topic_and_finish_actions(headers, course["id"], "Sorting")

        velocity = client.get("/velocity", headers=headers).json()

        assert velocity["actionsCompletedThisWeek"] == 3

    def test_velocity_is_zero_for_a_new_account(self) -> None:
        velocity = client.get("/velocity", headers=auth_headers()).json()

        assert velocity["actionsCompletedThisWeek"] == 0
        assert velocity["averageMastery"] is None


class TestConstraints:
    def test_constraints_round_trip(self) -> None:
        headers = auth_headers()
        payload = {
            "blockedSlots": [{"day": 2, "startTime": "09:00", "endTime": "17:00"}],
            "timePreference": "morning",
        }

        client.put("/constraints", json=payload, headers=headers)
        stored = client.get("/constraints", headers=headers).json()

        assert stored["timePreference"] == "morning"
        assert stored["blockedSlots"][0]["day"] == 2

    def test_defaults_before_anything_is_saved(self) -> None:
        stored = client.get("/constraints", headers=auth_headers()).json()

        assert stored["blockedSlots"] == []


def _make_pptx(*slide_titles: str) -> bytes:
    """A real, parseable .pptx in memory - not a stub, an actual deck."""
    from io import BytesIO

    from pptx import Presentation

    presentation = Presentation()
    layout = presentation.slide_layouts[0]
    for title in slide_titles:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


class TestUpload:
    @pytest.mark.parametrize("filename", ["notes.txt", "notes.docx"])
    def test_unsupported_file_types_are_rejected(self, filename: str) -> None:
        headers = auth_headers()
        course = create_course(headers)

        response = client.post(
            f"/courses/{course['id']}/topics/extract",
            files=[("files", (filename, b"some text", "text/plain"))],
            headers=headers,
        )

        assert response.status_code == 400

    def test_extracting_a_real_pptx_creates_topics_with_actions(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        content = _make_pptx("Binary Search Trees", "Hash Tables")

        response = client.post(
            f"/courses/{course['id']}/topics/extract",
            files=[("files", ("lecture1.pptx", content, "application/vnd.openxmlformats"))],
            headers=headers,
        )

        assert response.status_code == 201
        body = response.json()
        assert body["analysedBy"] == "heuristic"  # no AI key set for this student
        assert {topic["name"] for topic in body["created"]} == {
            "Binary Search Trees",
            "Hash Tables",
        }

        board = client.get(f"/board?courseId={course['id']}", headers=headers).json()
        actions_per_topic = {len(card["actions"]) for card in board["cards"]}
        assert actions_per_topic == {3}  # read, summarize, quiz - every topic, no exceptions

    def test_uploaded_file_lands_under_the_uploader_s_own_prefix(
        self, fake_storage
    ) -> None:
        headers = auth_headers("uploader@example.com")
        course = create_course(headers)
        user_id = client.get("/auth/me", headers=headers).json()["id"]

        client.post(
            f"/courses/{course['id']}/topics/extract",
            files=[("files", ("lecture1.pptx", _make_pptx("Recursion"), "application/x"))],
            headers=headers,
        )

        stored_keys = list(fake_storage.objects)
        assert len(stored_keys) == 1
        assert stored_keys[0].startswith(f"{user_id}/{course['id']}/")
        assert stored_keys[0].endswith("-lecture1.pptx")

    def test_a_different_students_upload_lands_under_their_own_prefix(
        self, fake_storage
    ) -> None:
        alice = auth_headers("alice@example.com")
        bob = auth_headers("bob@example.com")
        course = create_course(alice)
        client.post(
            f"/courses/{course['id']}/members", json={"email": "bob@example.com"}, headers=alice
        )

        client.post(
            f"/courses/{course['id']}/topics/extract",
            files=[("files", ("a.pptx", _make_pptx("Alice's slide"), "application/x"))],
            headers=alice,
        )
        client.post(
            f"/courses/{course['id']}/topics/extract",
            files=[("files", ("b.pptx", _make_pptx("Bob's slide"), "application/x"))],
            headers=bob,
        )

        alice_id = client.get("/auth/me", headers=alice).json()["id"]
        bob_id = client.get("/auth/me", headers=bob).json()["id"]
        prefixes = {key.split("/")[0] for key in fake_storage.objects}
        assert prefixes == {alice_id, bob_id}

    def test_too_many_files_are_rejected(self) -> None:
        headers = auth_headers()
        course = create_course(headers)

        response = client.post(
            f"/courses/{course['id']}/topics/extract",
            files=[
                ("files", (f"deck{index}.pdf", b"%PDF-1.4", "application/pdf"))
                for index in range(16)
            ],
            headers=headers,
        )

        assert response.status_code == 413


class TestSprintEndpoint:
    def test_empty_sprint_before_anything_is_planned(self) -> None:
        sprint = client.get("/sprint", headers=auth_headers()).json()

        assert sprint["status"] == "empty"
        assert sprint["committedMinutes"] == 0

    def test_pulling_a_topic_into_todo_commits_its_time(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        topic = client.post(
            f"/courses/{course['id']}/topics", json={"name": "Recursion"}, headers=headers
        ).json()

        client.patch(
            f"/topics/{topic['id']}/progress?courseId={course['id']}",
            json={"status": "todo"},
            headers=headers,
        )
        sprint = client.get("/sprint", headers=headers).json()

        # read 60 + summarize 45 + quiz 30
        assert sprint["committedMinutes"] == 135
        assert sprint["topicCount"] == 1

    def test_backlog_topics_are_not_part_of_the_commitment(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        client.post(f"/courses/{course['id']}/topics", json={"name": "Later"}, headers=headers)

        sprint = client.get("/sprint", headers=headers).json()

        assert sprint["committedMinutes"] == 0
        assert sprint["backlogCount"] == 1

    def test_over_committing_the_week_is_flagged(self) -> None:
        headers = auth_headers()
        course = create_course(headers)
        # Block every evening so capacity is near zero, then commit work anyway.
        client.put(
            "/constraints",
            json={
                "blockedSlots": [
                    {"day": day, "startTime": "15:00", "endTime": "23:00"} for day in range(7)
                ],
                "timePreference": "evening",
            },
            headers=headers,
        )
        topic = client.post(
            f"/courses/{course['id']}/topics", json={"name": "Recursion"}, headers=headers
        ).json()
        client.patch(
            f"/topics/{topic['id']}/progress?courseId={course['id']}",
            json={"status": "todo"},
            headers=headers,
        )

        sprint = client.get("/sprint", headers=headers).json()

        assert sprint["status"] in ("over_committed", "no_capacity")


def _add_topic_and_finish_actions(headers: dict[str, str], course_id: str, name: str) -> dict:
    """Create a topic and mark all three of its actions done. Returns the board card."""
    client.post(f"/courses/{course_id}/topics", json={"name": name}, headers=headers)
    board = client.get(f"/board?courseId={course_id}", headers=headers).json()

    for action in board["cards"][0]["actions"]:
        client.patch(
            f"/actions/{action['id']}/progress?courseId={course_id}",
            json={"isDone": True},
            headers=headers,
        )

    return client.get(f"/board?courseId={course_id}", headers=headers).json()["cards"][0]

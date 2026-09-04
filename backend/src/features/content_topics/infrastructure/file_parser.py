"""Pull raw text lines out of an uploaded PDF or PPTX.

Kept separate from the extraction heuristic so the algorithm can be tested with
plain lists of strings, with no file fixtures involved.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pypdf import PdfReader


class UnsupportedFileType(Exception):
    pass


def read_lines(filename: str, content: bytes) -> list[str]:
    """Document text as a flat list of lines, in reading order."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _read_pdf(content)
    if lower.endswith((".pptx", ".ppt")):
        return _read_pptx(content)

    raise UnsupportedFileType(filename)


def _read_pdf(content: bytes) -> list[str]:
    reader = PdfReader(io.BytesIO(content))
    lines: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        lines.extend(line for line in text.splitlines() if line.strip())
    return lines


def _read_pptx(content: bytes) -> list[str]:
    """Slide text, title first.

    A slide's title placeholder is almost always the topic, so it's emitted
    before the body text to keep the document order meaningful.
    """
    presentation = Presentation(io.BytesIO(content))
    lines: list[str] = []

    for slide in presentation.slides:
        title = getattr(slide.shapes, "title", None)
        if title is not None and title.has_text_frame and title.text.strip():
            lines.append(title.text.strip())

        for shape in slide.shapes:
            if shape is title or not shape.has_text_frame:
                continue
            lines.extend(
                paragraph.text.strip()
                for paragraph in shape.text_frame.paragraphs
                if paragraph.text.strip()
            )

    return lines
